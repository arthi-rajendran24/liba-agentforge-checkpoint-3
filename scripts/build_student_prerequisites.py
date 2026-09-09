from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "workshop" / "student-prerequisites-assets"
OUTPUT = ROOT / "workshop" / "AgentForge_Student_Prerequisites_Gemini_Telegram.docx"

BLACK = "000000"
INK = "17252A"
MUTED = "596A70"
TEAL = "0B8F83"
TEAL_DARK = "086E66"
PALE = "EAF7F5"
PALE_BLUE = "EEF5FB"
LIGHT = "F6F8F9"
WHITE = "FFFFFF"
GRAY = "D9E1E4"
RED = "B42318"
GREEN = "137A58"


def font_path(bold: bool = False) -> str:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/SFNS.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return path
    raise FileNotFoundError("No suitable system font found")


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(font_path(bold), size=size)


def rounded(draw, xy, radius=18, fill="#FFFFFF", outline=None, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def browser_frame(title: str, url: str):
    image = Image.new("RGB", (1600, 900), "#E9EEF0")
    draw = ImageDraw.Draw(image)
    rounded(draw, (40, 35, 1560, 865), 24, "#FFFFFF", "#CAD5D8", 2)
    draw.rounded_rectangle((40, 35, 1560, 125), radius=24, fill="#12242A")
    draw.rectangle((40, 85, 1560, 125), fill="#12242A")
    for x, color in [(78, "#FF6B63"), (112, "#F6C94D"), (146, "#61C77A")]:
        draw.ellipse((x - 10, 68, x + 10, 88), fill=color)
    draw.text((190, 57), title, font=pil_font(25, True), fill="#F2F7F8")
    rounded(draw, (235, 92, 1365, 115), 10, "#E8EFF1")
    draw.text((260, 92), url, font=pil_font(16), fill="#485E66")
    return image, draw


def step_badge(draw, x, y, number, text, width=390):
    draw.ellipse((x, y, x + 52, y + 52), fill="#0B8F83")
    n = str(number)
    box = draw.textbbox((0, 0), n, font=pil_font(24, True))
    draw.text((x + 26 - (box[2] - box[0]) / 2, y + 11), n, font=pil_font(24, True), fill="#FFFFFF")
    draw.text((x + 68, y + 8), text, font=pil_font(25, True), fill="#17252A")
    draw.line((x + 68, y + 43, x + width, y + 43), fill="#C9D6D8", width=2)


def save_screen_maps():
    ASSETS.mkdir(parents=True, exist_ok=True)

    image, draw = browser_frame("Google AI Studio", "aistudio.google.com/apikey")
    draw.text((95, 165), "Gemini API keys", font=pil_font(42, True), fill="#17252A")
    draw.text((95, 222), "Create and manage keys for the Gemini API", font=pil_font(24), fill="#596A70")
    rounded(draw, (1230, 170, 1495, 230), 14, "#0B8F83")
    draw.text((1270, 185), "+  Create API key", font=pil_font(23, True), fill="#FFFFFF")
    rounded(draw, (95, 285, 1495, 650), 16, "#F7F9FA", "#D7E0E3")
    draw.text((130, 320), "Project", font=pil_font(21, True), fill="#53666E")
    draw.text((535, 320), "Key name", font=pil_font(21, True), fill="#53666E")
    draw.text((1055, 320), "Key type", font=pil_font(21, True), fill="#53666E")
    draw.line((125, 365, 1465, 365), fill="#D7E0E3", width=2)
    draw.text((130, 405), "My first project", font=pil_font(25), fill="#17252A")
    draw.text((535, 405), "AgentForge workshop", font=pil_font(25), fill="#17252A")
    draw.text((1055, 405), "Authorization key", font=pil_font(25), fill="#137A58")
    draw.text((95, 710), "Click this button after signing in", font=pil_font(24, True), fill="#0B8F83")
    draw.line((1210, 695, 1350, 235), fill="#0B8F83", width=6)
    draw.polygon([(1350, 235), (1326, 257), (1362, 265)], fill="#0B8F83")
    image.save(ASSETS / "01-gemini-api-key-screen-map.png")

    image, draw = browser_frame("Google AI Studio", "aistudio.google.com/apikey")
    rounded(draw, (390, 175, 1210, 755), 22, "#FFFFFF", "#AFC2C7", 3)
    draw.text((445, 225), "Create API key", font=pil_font(38, True), fill="#17252A")
    draw.text((445, 290), "Select a project", font=pil_font(23, True), fill="#53666E")
    rounded(draw, (445, 330, 1155, 405), 12, "#F7F9FA", "#CAD5D8")
    draw.text((475, 351), "My first project", font=pil_font(25), fill="#17252A")
    draw.text((445, 450), "Key name", font=pil_font(23, True), fill="#53666E")
    rounded(draw, (445, 490, 1155, 565), 12, "#FFFFFF", "#CAD5D8")
    draw.text((475, 511), "AgentForge workshop", font=pil_font(25), fill="#17252A")
    rounded(draw, (835, 635, 1155, 700), 14, "#0B8F83")
    draw.text((890, 653), "Create API key", font=pil_font(25, True), fill="#FFFFFF")
    step_badge(draw, 95, 235, 1, "Choose a project")
    step_badge(draw, 95, 430, 2, "Use a clear name")
    step_badge(draw, 95, 625, 3, "Create and copy once")
    image.save(ASSETS / "02-gemini-create-key-screen-map.png")

    image, draw = browser_frame("Telegram Applications", "telegram.org/apps")
    draw.text((95, 165), "Install the official Telegram app", font=pil_font(42, True), fill="#17252A")
    draw.text((95, 222), "Use your personal phone. Desktop installation on the lab computer is not required.", font=pil_font(23), fill="#596A70")
    cards = [
        (95, 305, "Android", "Google Play or official Android app"),
        (575, 305, "iPhone and iPad", "Apple App Store"),
        (1055, 305, "Desktop", "Optional for your own computer"),
    ]
    for x, y, title, sub in cards:
        rounded(draw, (x, y, x + 410, y + 285), 20, "#F7F9FA", "#CAD5D8")
        draw.ellipse((x + 145, y + 35, x + 265, y + 155), fill="#2AABEE")
        draw.polygon([(x + 172, y + 98), (x + 242, y + 66), (x + 218, y + 130), (x + 200, y + 111)], fill="#FFFFFF")
        draw.text((x + 30, y + 180), title, font=pil_font(28, True), fill="#17252A")
        draw.multiline_text((x + 30, y + 225), sub, font=pil_font(19), fill="#596A70", spacing=6)
    draw.text((95, 665), "After installation", font=pil_font(27, True), fill="#0B8F83")
    draw.text((95, 712), "Sign in, confirm you can receive messages, and keep your phone with you for the workshop.", font=pil_font(24), fill="#17252A")
    image.save(ASSETS / "03-telegram-install-screen-map.png")

    image = Image.new("RGB", (1600, 980), "#DCE9EE")
    draw = ImageDraw.Draw(image)
    rounded(draw, (160, 40, 1440, 940), 30, "#F7FAFB", "#BAC9CE", 3)
    draw.rectangle((160, 40, 1440, 150), fill="#17252A")
    draw.ellipse((205, 67, 275, 137), fill="#2AABEE")
    draw.text((305, 68), "BotFather", font=pil_font(32, True), fill="#FFFFFF")
    draw.text((305, 108), "@BotFather  •  official Telegram bot", font=pil_font(19), fill="#B9D7E1")
    rounded(draw, (760, 205, 1370, 285), 24, "#DCF8C6")
    draw.text((800, 228), "/newbot", font=pil_font(27, True), fill="#17252A")
    rounded(draw, (230, 320, 1130, 430), 24, "#FFFFFF", "#D7E0E3")
    draw.text((270, 345), "Alright, a new bot. What are we going to call it?", font=pil_font(24), fill="#17252A")
    draw.text((270, 385), "Please choose a display name.", font=pil_font(22), fill="#596A70")
    rounded(draw, (760, 465, 1370, 555), 24, "#DCF8C6")
    draw.text((800, 493), "LIBA Team 12 AgentForge", font=pil_font(25, True), fill="#17252A")
    rounded(draw, (230, 590, 1250, 720), 24, "#FFFFFF", "#D7E0E3")
    draw.text((270, 615), "Now choose a username for your bot.", font=pil_font(24), fill="#17252A")
    draw.text((270, 657), "It must be unique and end in bot.", font=pil_font(22), fill="#596A70")
    rounded(draw, (680, 760, 1370, 850), 24, "#DCF8C6")
    draw.text((720, 788), "liba_team12_agentforge_bot", font=pil_font(24, True), fill="#17252A")
    draw.text((210, 880), "Illustration only  •  Names shown are examples", font=pil_font(19), fill="#596A70")
    image.save(ASSETS / "04-botfather-newbot-screen-map.png")

    image = Image.new("RGB", (1600, 900), "#F3F7F8")
    draw = ImageDraw.Draw(image)
    draw.text((90, 70), "What to record after BotFather creates the bot", font=pil_font(42, True), fill="#17252A")
    draw.text((90, 130), "Store this privately. Never send it in the class group or place it in a screenshot.", font=pil_font(23), fill="#596A70")
    fields = [
        ("Bot display name", "LIBA Team 12 AgentForge"),
        ("Bot username", "@liba_team12_agentforge_bot"),
        ("Numeric bot ID", "1234567890"),
        ("Bot API token", "1234567890:••••••••••••••••••••••••"),
    ]
    y = 220
    for label, value in fields:
        draw.text((120, y), label, font=pil_font(22, True), fill="#53666E")
        rounded(draw, (120, y + 38, 1480, y + 112), 12, "#FFFFFF", "#C8D5D9")
        draw.text((150, y + 58), value, font=pil_font(26, True if "token" in label.lower() else False), fill="#17252A")
        y += 145
    draw.text((120, 800), "The number before the first colon in the standard token is the bot ID.", font=pil_font(22, True), fill="#0B8F83")
    image.save(ASSETS / "05-telegram-credential-record.png")


def set_run_font(run, name="Aptos"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=130, start=150, bottom=130, end=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
    no_split = OxmlElement("w:cantSplit")
    no_split.set(qn("w:val"), "true")
    tr_pr.append(no_split)


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    no_split.set(qn("w:val"), "true")
    tr_pr.append(no_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text, url, color=TEAL_DARK):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color_node, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(34)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_style.paragraph_format.space_after = Pt(12)
    title_ppr = title_style._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

    for style_name, size in (("Heading 1", 22), ("Heading 2", 15), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "AGENTFORGE  |  STUDENT PREPARATION"
    for run in hp.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BLACK)

    footer = section.footer
    footer.paragraphs[0].text = ""
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.8))
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer_table.columns[0].width = Inches(5.3)
    footer_table.columns[1].width = Inches(1.5)
    left_footer = footer_table.cell(0, 0).paragraphs[0]
    left_footer.text = "LIBA workshop  •  Complete before the session"
    for run in left_footer.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    right_footer = footer_table.cell(0, 1).paragraphs[0]
    add_page_number(right_footer)
    for run in right_footer.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    def para(text="", *, bold=False, size=None, color=None, align=None, italic=False, keep=False):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.keep_with_next = keep
        r = p.add_run(text)
        set_run_font(r)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def bullet(text, level=0):
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_run_font(r)
        return p

    def numbered(number, text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.32)
        p.paragraph_format.space_after = Pt(6)
        nr = p.add_run(f"{number}.  ")
        set_run_font(nr)
        nr.bold = True
        nr.font.color.rgb = RGBColor.from_string(TEAL_DARK)
        r = p.add_run(text)
        set_run_font(r)
        return p

    def rule():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), TEAL)
        borders.append(bottom)
        p_pr.append(borders)

    def figure(filename, caption, width=6.8):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        path = ASSETS / filename
        pic = p.add_run().add_picture(str(path), width=Inches(width))
        pic._inline.docPr.set("descr", caption)
        pic._inline.docPr.set("title", path.stem.replace("-", " ").title())
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        for run in cp.runs:
            set_run_font(run)
            run.font.size = Pt(8.5)
            run.italic = True
            run.font.color.rgb = RGBColor.from_string(MUTED)

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        set_table_header(t.rows[0])
        for idx, header in enumerate(headers):
            cell = t.rows[0].cells[idx]
            cell.text = header
            set_cell_shading(cell, "24434D")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
                    run.font.size = Pt(9.5)
        for row_idx, row in enumerate(rows):
            tr = t.add_row()
            keep_row(tr)
            for idx, value in enumerate(row):
                cell = tr.cells[idx]
                cell.text = str(value)
                cell.width = Inches(widths[idx])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if row_idx % 2:
                    set_cell_shading(cell, PALE_BLUE)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        set_run_font(run)
                        run.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def checkbox(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"☐  {text}")
        set_run_font(r)
        return p

    def page_break():
        doc.add_page_break()

    para("STUDENT CHECKLIST  |  SEPTEMBER 2026", bold=True, color=TEAL_DARK, size=10)
    doc.add_paragraph("AgentForge Student Prerequisites", style="Title")
    para("Gemini API and Telegram bot preparation for the LIBA workshop", size=17, color=BLACK)
    para("Please complete these two items before the session", bold=True, size=18)
    para(
        "You will use a lab computer during the workshop. The institute technology team will prepare the machine with Antigravity, Git, uv and the remaining software. You only need to arrive with access to your own Gemini API key and a Telegram bot that you created in advance."
    )
    para("Bring your phone", bold=True, size=14)
    para("Keep your phone with you, signed in to Telegram and able to access your Google account. You may need your normal login or two step verification during the session.")
    para("Do not send either credential in advance", bold=True, size=14)
    para("Your Gemini key and Telegram bot token work like passwords. Store them privately and enter them only when I guide you during the workshop.")
    para("Time required", bold=True, size=14)
    para("Allow about 15 to 20 minutes. Complete the setup at least one day before the workshop so there is time to solve account or username issues.")

    page_break()
    doc.add_heading("Your two prerequisites", level=1)
    para("Finish both columns. Nothing else needs to be installed by you on the lab computer.")
    table(
        ["Gemini API", "Telegram bot"],
        [
            ["Sign in to Google AI Studio", "Install the official Telegram app on your phone"],
            ["Create or locate one Gemini API key", "Sign in and open the official @BotFather"],
            ["Store the key privately", "Create one bot with /newbot"],
            ["Do not add billing unless told to", "Record the bot username, numeric ID and token privately"],
        ],
        [3.35, 3.35],
    )
    doc.add_heading("What you need with you", level=2)
    bullet("Your personal phone with Telegram installed and signed in")
    bullet("Access to the Google account used for Google AI Studio")
    bullet("Access to any verification method required by Google or Telegram")
    bullet("A private password manager or secure note that only you can open")
    doc.add_heading("What the institute technology team handles", level=2)
    para("The lab team will prepare Antigravity, Git, uv, the AgentForge repository and all other machine-level software. Please do not spend time installing those tools on the lab computers before the session.")
    doc.add_heading("Credential rule", level=2)
    para("Do not paste your key or token into WhatsApp, email, a class group, a shared document, GitHub, a screenshot or a public AI chat. During the workshop, you will enter them locally when I ask you to.", bold=True)

    page_break()
    doc.add_heading("Part 1 Create your Gemini API key", level=1)
    para("Google AI Studio is the official place to create and manage a Gemini API key.")
    p = doc.add_paragraph()
    p.add_run("Open: ").bold = True
    add_hyperlink(p, "Google AI Studio API Keys page", "https://aistudio.google.com/apikey")
    p.add_run("  aistudio.google.com/apikey")
    numbered(1, "Open the link and sign in with the Google account you will be able to access during the workshop.")
    numbered(2, "Accept the Google AI Studio terms if this is your first visit.")
    numbered(3, "Open the API Keys page. New users may already see a default project and key.")
    numbered(4, "If you need a new key, select Create API key.")
    figure("01-gemini-api-key-screen-map.png", "Illustrated screen guide based on the Google AI Studio API Keys page. Button placement may vary slightly.", 6.65)

    page_break()
    doc.add_heading("Create and store the Gemini key", level=1)
    numbered(5, "Choose your default project or another personal project that you control. If an institutional account blocks key creation, try your personal Google account or ask the workshop coordinator for help.")
    numbered(6, "Give the key a clear name such as AgentForge workshop.")
    numbered(7, "Select Create API key. Google currently creates new AI Studio keys as authorization keys.")
    numbered(8, "Copy the key once and save it in your password manager or a private note. Do not include it in the document you submit for the workshop.")
    figure("02-gemini-create-key-screen-map.png", "Illustrated screen guide for the key creation dialog. Never place a real key in a screenshot.", 6.65)
    doc.add_heading("Billing", level=2)
    para("A paid subscription is not a student prerequisite for this workshop. Do not add a payment method or upgrade the project unless I or the institute gives you a separate instruction.")
    doc.add_heading("Your Gemini readiness check", level=2)
    checkbox("I can sign in to Google AI Studio.")
    checkbox("I can see a Gemini API key in the API Keys page.")
    checkbox("The key is stored privately and I can retrieve it during the session.")
    checkbox("I have not shared, photographed or committed the key anywhere.")

    page_break()
    doc.add_heading("Gemini problems you can solve before class", level=1)
    table(
        ["What you see", "What to do"],
        [
            ["Create API key is unavailable", "Use a personal Google account or ask the owner of the Cloud project for permission."],
            ["AI Studio asks you to accept terms", "Read and accept the terms using the account you intend to use."],
            ["Your institute account is restricted", "Use a personal account if permitted, or tell the workshop coordinator before the session."],
            ["A key is marked blocked", "Create a new key in AI Studio. Do not reuse a blocked or exposed key."],
            ["You accidentally shared a key", "Treat it as compromised. Generate a replacement and disable or revoke the exposed key."],
            ["AI Studio asks about billing", "Stay on the free tier for this workshop unless the institute gives a separate instruction."],
        ],
        [2.3, 4.4],
    )
    doc.add_heading("Why the key must remain private", level=2)
    para("Anyone who has the key may consume your project quota or create costs if billing is enabled. Google recommends keeping keys out of source control and client-side applications. AgentForge will read the key from a protected local environment when we connect it together.")
    doc.add_heading("Do not test by posting the key", level=2)
    para("You do not need to paste the key into a website, a public API tester or a chat to prove that it exists. Seeing the key in your AI Studio account and storing it securely is enough for the prerequisite.")

    page_break()
    doc.add_heading("Part 2 Install Telegram", level=1)
    para("Install the official Telegram app on your personal phone and sign in before creating the bot.")
    p = doc.add_paragraph()
    p.add_run("Official downloads: ").bold = True
    add_hyperlink(p, "Telegram official apps page", "https://telegram.org/apps")
    p.add_run("  telegram.org/apps")
    numbered(1, "Install Telegram from the official Telegram page, Google Play or the Apple App Store.")
    numbered(2, "Sign in with the phone number you normally use for Telegram and complete Telegram's verification steps.")
    numbered(3, "Confirm that you can open a chat and receive a message. Keep the app signed in for the workshop.")
    figure("03-telegram-install-screen-map.png", "Illustrated screen guide for Telegram's official app choices.", 6.65)
    doc.add_heading("Use your phone", level=2)
    para("Telegram Desktop is optional on your personal computer. You do not need to install or sign in to Telegram on the shared lab machine before class.")

    page_break()
    doc.add_heading("Create your Telegram bot with BotFather", level=1)
    p = doc.add_paragraph()
    p.add_run("Open the official bot: ").bold = True
    add_hyperlink(p, "Official BotFather chat", "https://t.me/BotFather")
    p.add_run("  t.me/BotFather")
    numbered(1, "In Telegram, search for @BotFather. Confirm the exact handle and the official verification mark before continuing.")
    numbered(2, "Select Start, then send /newbot.")
    numbered(3, "Choose a display name. A simple format is LIBA Team followed by your team number and AgentForge.")
    numbered(4, "Choose a unique username that ends in bot. Example: liba_team12_agentforge_bot. If it is taken, add your section or another number.")
    figure("04-botfather-newbot-screen-map.png", "Illustrated BotFather conversation. The example name and username are fictional.", 6.35)
    para("The bot may not reply to messages yet. That is expected. We will connect it to AgentForge together during the session.", bold=True)

    page_break()
    doc.add_heading("Record the Telegram bot details", level=1)
    para("After the username is accepted, BotFather sends the bot's API token. Copy the details below into a password manager or private note that only you can access.")
    figure("05-telegram-credential-record.png", "Illustrated credential record with a masked example token. Do not create a screenshot containing your real token.", 6.5)
    doc.add_heading("The four details to keep", level=2)
    table(
        ["Detail", "Example", "Where it comes from"],
        [
            ["Display name", "LIBA Team 12 AgentForge", "The first name you give BotFather"],
            ["Bot username", "@liba_team12_agentforge_bot", "The unique username accepted by BotFather"],
            ["Numeric bot ID", "1234567890", "The digits before the first colon in the standard bot token"],
            ["Bot API token", "1234567890:••••••", "The full private token returned by BotFather"],
        ],
        [1.35, 2.4, 2.95],
    )
    page_break()
    doc.add_heading("Telegram safety and readiness", level=1)
    doc.add_heading("Do not create a Telegram developer application", level=2)
    para("AgentForge will use the Bot API token issued by BotFather. Students do not need a separate api_id or api_hash from my.telegram.org. The numeric bot ID can be verified automatically when we connect the token during the workshop.")
    doc.add_heading("If the token is exposed", level=2)
    para("Open BotFather and use its bot-management controls to revoke or regenerate the token. Save the replacement privately. Anyone with the token can control the bot.")
    doc.add_heading("Before class", level=2)
    checkbox("I can open the chat with the official @BotFather.")
    checkbox("I can open my new bot's Telegram profile using its username.")
    checkbox("My bot username is unique and ends in bot.")
    checkbox("My complete token is available in private storage.")
    checkbox("I have not posted the token or placed it in a screenshot.")
    doc.add_heading("A silent bot is normal", level=2)
    para("Creating a bot registers its identity and gives you the token. It does not give the bot AgentForge behavior yet. If you open the bot and it does not respond, your prerequisite can still be complete. We will connect and test the behavior together.")

    page_break()
    doc.add_heading("Final student readiness checklist", level=1)
    para("Complete this page at least one day before the workshop. Keep the credential values in your private storage; write only the non-secret details here.")
    doc.add_heading("Gemini", level=2)
    checkbox("I can sign in to Google AI Studio with the account I will use in class.")
    checkbox("I have one Gemini API key ready.")
    checkbox("The key is stored privately and has never been posted or shared.")
    doc.add_heading("Telegram", level=2)
    checkbox("The official Telegram app is installed and signed in on my phone.")
    checkbox("I created my bot using the official @BotFather.")
    checkbox("My bot username ends in bot and I can open its profile.")
    checkbox("I stored the bot username, numeric ID and complete token privately.")
    checkbox("I understand that the bot may stay silent until we connect it in class.")
    doc.add_heading("Bring to class", level=2)
    checkbox("My phone is charged and I can receive Google or Telegram verification prompts.")
    checkbox("I can open my private credential storage without asking someone else for access.")
    para("Student name  __________________________________________", size=11)
    para("Section and team  ________________________________________", size=11)
    para("Telegram bot username only  @______________________________", size=11)
    para("Do not write your Gemini key or Telegram token on this page.", bold=True, color=RED)

    page_break()
    doc.add_heading("Send this message to students", level=1)
    para("You can copy and send this section with the document.")
    para(
        "Hi everyone, before our AgentForge workshop, please complete two things: first, create a Gemini API key in Google AI Studio; second, install Telegram on your phone and create one bot using the official @BotFather. Keep both the Gemini key and Telegram bot token private. Do not send them to me or post them in the class group. Bring your phone and make sure you can access the Google account and Telegram account you used. The institute technology team will prepare Antigravity, Git, uv, AgentForge and the other software on the lab machines. Please finish the attached checklist at least one day before the session. We will connect your bot to AgentForge together during the workshop."
    )
    doc.add_heading("The confirmation students can send", level=2)
    para("Done. I have created my Gemini API key and Telegram bot, and I have stored both credentials privately. My bot username is @______________________. I will bring my phone and will not send either secret in this message.")
    para("Students should send only this confirmation and the bot username. They should never include the Gemini key or Telegram token.", bold=True)
    doc.add_heading("Official references", level=2)
    sources = [
        ("Google AI Studio API Keys", "https://aistudio.google.com/apikey"),
        ("Google Gemini API key documentation", "https://ai.google.dev/gemini-api/docs/api-key"),
        ("Telegram official applications", "https://telegram.org/apps"),
        ("Telegram BotFather tutorial", "https://core.telegram.org/bots/tutorial"),
        ("Official BotFather", "https://t.me/BotFather"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)
    para("Instructions and interface labels checked on 4 September 2026. Screens may change slightly by device or account.", italic=True, color=MUTED, size=9.5)

    doc.core_properties.title = "AgentForge Student Prerequisites"
    doc.core_properties.subject = "Gemini API and Telegram bot preparation for LIBA students"
    doc.core_properties.author = "Arthi"
    doc.core_properties.keywords = "AgentForge, LIBA, Gemini API, Telegram, BotFather, student prerequisites"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    save_screen_maps()
    print(build_document())
